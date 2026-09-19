"""Match historical baseline errors to the ensemble's fixed evaluation rows."""
import csv,math
import numpy as np
from openpyxl import load_workbook
from paper_scope import HISTORICAL_COUNT, PDE_COUNT, PAPER_COUNT
from elo_ranking import rank_errors
from build_data import ROOT,NAMES,CLASS_OF,EXCLUSIONS,load,dump,digest,table,tex,ordered_datasets

MODEL_INFO={'mf_fno_transfer_film': ('M1',
                          'FiLM FNO transfer',
                          'li2021fno,perez2018',
                          'Coordinate input and parameter FiLM at every FNO block, with LF to HF transfer. '
                          'Tests repeated conditioning and spectral transfer.'),
 'mf_fno_allpairs': ('M2',
                     'All pairs FNO',
                     'li2021fno,herde2024',
                     'Fidelity labels and repeated target supervision. Poseidon is pair enumeration '
                     'inspiration only. No source field input.'),
 'convnext_unet_film': ('M3',
                        'ConvNeXt transfer',
                        'liu2022convnext',
                        'ConvNeXt blocks in a U shaped field decoder, parameter FiLM, and LF to HF '
                        'transfer. Tests local spatial processing.'),
 'fno_fire_distcond': ('M4',
                       'Distribution FNO',
                       'yu2026fire',
                       'FIRE distribution summaries become spatial channels from five LF FNOs. A trained '
                       'residual FNO replaces foundation model inference.'),
 'wno_transfer_film': ('M5',
                       'Wavelet transfer',
                       'tripura2022wno',
                       'WNO motivates a custom single level db4 coefficient mixer with FiLM and transfer '
                       'training.'),
 'mf_deeponet': ('M6',
                 'Retrieved field DeepONet',
                 'lu2022mf',
                 'Parameters and a nearest neighbour LF training field enter the branch. The residual '
                 'needs no new coarse solve.'),
 'st_hf_pod_gp': ('M7',
                  'POD GP',
                  'higdon2008',
                  'HF field SVD and shared coefficient GP hyperparameters. Basis emulation inspiration, not '
                  'the original Bayesian model.'),
 'uqcorr_transolver_pred': ('M8',
                            'Slice attention corrector',
                            'wu2024transolver',
                            'Changed pooling and cell to slice attention, parameter FiLM, and six '
                            'corrections of a predicted coarse field.'),
 'uqcorr_convnext_pred': ('M9',
                          'ConvNeXt corrector',
                          'liu2022convnext',
                          'ConvNeXt blocks in a U shaped corrector with FiLM. Six updates refine a '
                          'predicted coarse field rather than generating it directly.'),
 'st_koh_pod': ('B1',
                'Autoregressive POD GP',
                'kennedy2000',
                'KOH inspired scalar field coupling plus a separate residual POD GP. Not joint probabilistic '
                'co kriging.'),
 'st_nargp_pod': ('B2',
                  'Nonlinear POD GP',
                  'perdikaris2017',
                  'NARGP inspired coefficient kernel with shared covariance hyperparameters and Monte Carlo '
                  'LF uncertainty.'),
 'st_mfdnn': ('B3',
              'Composite MF MLP',
              'meng2020',
              'Meng inspired pointwise LF network and staged linear plus nonlinear HF corrections.'),
 'st_mfdeeponet': ('B4',
                   'Composite MF DeepONet',
                   'howard2023',
                   'Howard inspired LF sensor prediction and staged linear plus nonlinear DeepONets.'),
 'st_dmfal': ('B5',
              'Latent MF network',
              'li2022dmfal',
              'DMFAL inspires two latent levels. Deterministic MSE replaces the probabilistic objective and '
              'acquisition.'),
 'nomad_mf': ('B6',
              'Nonlinear decoder transfer',
              'seidman2022',
              'NOMAD motivates a custom nonlinear coordinate decoder. Multifidelity behavior comes from '
              'transfer training.'),
 'mfrnp': ('B7',
           'MFRNP adapter',
           'niu2024',
           'Upstream MFRNP model with data, normalization, resize, batching, context, and training '
           'adaptations.'),
 'fno_coregionalization': ('B8',
                           'Fidelity basis FNO',
                           'li2022ifc',
                           'IFC inspires fidelity dependence. Spatial basis FNO omits the original neural '
                           'ODE and GP.'),
 'mf_fno_transfer': ('B9',
                     'FNO transfer I',
                     'lyu2023',
                     'Parameters concatenated with coordinates and LF to HF transfer. Same inspected recipe '
                     'as B10.'),
 'mf_fno_transfer_bar': ('B10',
                         'FNO transfer II',
                         'lyu2023',
                         'Separately archived transfer run with the same inspected model and training code '
                         'as B9.'),
 'st_lf_affine_pod': ('B11',
                      'Affine POD GP',
                      'zhang2018',
                      'One affine slope and intercept across fields applied to an LF POD GP. Conceptual '
                      'inspiration only.'),
 'st_knn': ('B12', 'Parameter kNN', '', 'Training-field nearest-neighbor reference.'),
 'st_mean': ('B13', 'Training mean', '', 'Input-independent fine-training mean.'),
 'st_film_hf_only': ('C1',
                     'FiLM HF-only',
                     '',
                     'Control. archived case identity not independently established.'),
 'st_film_lf_only': ('C2',
                     'FiLM LF-only',
                     '',
                     'Control. archived case identity not independently established.'),
 'st_film_pooled': ('C3',
                    'FiLM pooled',
                    '',
                    'Control. archived case identity not independently established.'),
 'st_film_affine': ('C4',
                    'FiLM affine',
                    '',
                    'Control. archived case identity not independently established.')}
RULES={'single_l2':('Selected model',''), 'uniform':('Uniform mixture',''),
       'inverse_mse':('Inverse error mixture','goel2007'), 'full':('Fitted mixture','breiman1996')}

def agg(errors,ref):
    ds=sorted(errors);lr={d:math.log(errors[d]/ref[d]) for d in ds};cs=sorted({CLASS_OF[d] for d in ds})
    return dict(n=len(ds),class_ratio=math.exp(np.mean([np.mean([lr[d] for d in ds if CLASS_OF[d]==c]) for c in cs])),
                dataset_ratio=math.exp(np.mean(list(lr.values()))),wins=sum(v<math.log(.99) for v in lr.values()),
                losses=sum(v>math.log(1.01) for v in lr.values()))

def fmt(value):
    s=f'{100*value:.3g}'
    if 'e' in s:
        a,b=s.split('e');s=a+'e'+str(int(b))
    return s

def main():
    from verify_b8_recovery import verify
    verify()
    base=load(ROOT/'data/baseline_alignment.json');remote={r['dataset']:r for r in base['remote_checks']}
    individual=load(ROOT/'data/individual_model_summary.json');s=load(ROOT/'data/paper_primary_summary.json')['5']
    library=individual['checks']['models'];means={d:dict(v,uniform=s['mean_errors'][d]['uniform']) for d,v in individual['means'].items()}
    records=[];missing=[];proofs=[]
    for r in base['raw_sources']:
        d=r['dataset'];m=r['model']
        if d not in means:continue
        rr=remote[d]
        if r.get('missing'):missing.append(dict(dataset=d,model=m,reason='No saved per-case errors'));continue
        assert digest(ROOT/'data'/r['file'])==r['sha256']
        raw=load(ROOT/'data'/r['file']);errors=np.asarray(raw['splits']['test_hf']['rel_l2_per_sample'])
        assert len(errors)==rr['n'] and list(raw['work_grid'])==rr['grid']
        assert np.isfinite(errors).all() and (errors>=0).all()
        assert rr['input_order_matches_reference']
        hashes=raw.get('manifest',{}).get('sha256',{})
        hashmatch=any(v==rr['test_archive_sha256'] and '/test_l' in k for k,v in hashes.items())
        arc=rr['archives'].get(m)
        exportmatch=bool(arc and arc['target_rows_match'] and np.allclose(errors,arc['errors'],rtol=1e-5,atol=1e-9))
        if not(hashmatch or exportmatch):
            missing.append(dict(dataset=d,model=m,reason='No source-file hash or matching saved target/error export'));continue
        proof=('source HF file hash' if hashmatch else
               'original checkpoint replay on identified test rows' if arc.get('checkpoint_replay_verified') else
               'saved target rows and pre-export errors')
        selected=[]
        for seed in [71,172,273]:
            split=load(ROOT/'data/historical_runs'/f'{d}__s{seed}'/'split.json')
            rows=[v['row'] for v in split['evaluation']]
            for v in split['evaluation']:
                assert v['group'].rsplit(':',1)[1]==rr['input_row_sha256'][v['row']]
            for cal in split['calibration'].values():assert not set(rows)&{v['row'] for v in cal}
            values=errors[rows];selected.append(float(values.mean()))
            for row,value in zip(rows,values):records.append(dict(dataset=d,model=m,base_seed=42,partition_seed=seed,row=row,relative_l2=float(value)))
        means[d][m]=float(np.mean(selected));proofs.append(dict(dataset=d,model=m,proof=proof,source=r['file'],seed=42))
    from extension_results import add_baselines
    add_baselines(means,records,proofs)
    numerator={};denominator={}
    for ds in individual['means']:
        parts=[r['errors'] for r in individual['partitions'] if r['dataset']==ds]
        numerator[ds]=float(np.mean([r['full'] for r in parts]))
        denominator[ds]=float(np.mean([min(r[m] for m in library) for r in parts]))
    dump(ROOT/'data/partition_hindsight_summary.json',agg(numerator,denominator))
    stats={};ref={d:v['single_l2'] for d,v in means.items()}
    for m in [*MODEL_INFO,*RULES]:
        e={d:v[m] for d,v in means.items() if m in v}
        if e:stats[m]=agg(e,ref)
    order=ordered_datasets(means)
    # Retrospective minima describe the available library, never deployment selection.
    paper_models=[m for m,v in MODEL_INFO.items() if v[0] in [f'B{i}' for i in range(1,12)]]
    best_library={d:min(means[d][m] for m in library) for d in order}
    best_baseline={d:min(means[d][m] for m in paper_models if m in means[d]) for d in order}
    headline_specs=[
        ('Best library / best additional baseline',best_library,best_baseline),
        ('Full / best library',{d:means[d]['full'] for d in order},best_library),
        ('Inverse / best library',{d:means[d]['inverse_mse'] for d in order},best_library),
        ('Full / best additional baseline',{d:means[d]['full'] for d in order},best_baseline),
        ('Full / selected model',{d:means[d]['full'] for d in order},ref)]
    headlines={name:agg(a,b) for name,a,b in headline_specs}
    dump(ROOT/'data/headline_comparisons.json',dict(comparisons=headlines,
         definition='Best means the minimum partition averaged error on each dataset. Retrospective diagnostics only.',
         additional_baselines=paper_models,
         missing_IFC_cells=sum('fno_coregionalization' not in means[d] for d in order)))
    table('headline_comparisons.tex','lrrrr',
          'Comparison & Class ratio & Dataset ratio & Wins & Losses',
          [[name,f"{v['class_ratio']:.3f}",f"{v['dataset_ratio']:.3f}",v['wins'],v['losses']] for name,v in headlines.items()])
    # Full row-level record permits independent recalculation of every added mean.
    for name,rows in [('matched_baseline_cases.csv',records),('matched_baseline_coverage.csv',proofs)]:
        positions={d:i for i,d in enumerate(order)}
        rows.sort(key=lambda r:positions[r['dataset']])
        with (ROOT/'data'/name).open('w',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    main_rules=['single_l2','inverse_mse','full']
    methods=library+main_rules
    rows=[]
    for d in order:
        baseline=min((m for m in paper_models if m in means[d]),key=lambda m:means[d][m])
        vals=[means[d][baseline]]+[means[d][m] for m in methods];best=min(vals)
        cells=[r'\textbf{'+fmt(v)+'}' if v==best else fmt(v) for v in vals]
        cells[0]+=' ('+MODEL_INFO[baseline][0]+')'
        rows.append([tex(NAMES[d])+(r'$^{\dagger}$' if d in EXCLUSIONS else ''),*cells])
    era=load(ROOT/'data/era5_summary.json')['means']
    era_baselines=load(ROOT/'data/era5_baselines_summary.json')['means']
    era_best_model=min((m for m in paper_models if m in era_baselines),key=lambda m:era_baselines[m])
    era_vals=[era_baselines[era_best_model]]+[era['expert:'+m] for m in library]+[era['b5__'+m] for m in ['selected_single','inverse_mse','full']]
    era_best=min(v for v in era_vals if v is not None)
    era_cells=['---' if v is None else (r'\textbf{'+fmt(v)+'}' if v==era_best else fmt(v)) for v in era_vals]
    era_cells[0]+=' ('+MODEL_INFO[era_best_model][0]+')'
    rows.append([r'\midrule ERA5',*era_cells])
    table('all_experts_main.tex','l'+'r'+'|'+'r'*9+'|'+'r'*3,
          'Dataset & Best baseline & '+' & '.join([f'M{i}' for i in range(1,10)]+['Selected','Inverse','Fitted']),rows)
    main_models=library+[m for m,v in MODEL_INFO.items() if v[0] in [f'B{i}' for i in range(1,12)]]
    # Table 2 covers all 22 datasets for its dataset ratio and Elo ranking.
    # The class ratio keeps the seven PDE classes and excludes climate emulation.
    ranked_methods=main_models+main_rules
    complete_means={d:{m:means[d][m] for m in ranked_methods} for d in order}
    complete_means['era5']={m:era['expert:'+m] for m in library}
    complete_means['era5'].update({m:era_baselines[m] for m in paper_models})
    complete_means['era5'].update({m:era['b5__'+('selected_single' if m=='single_l2' else m)] for m in main_rules})
    assert len(complete_means)==PAPER_COUNT
    elo=rank_errors(complete_means,ranked_methods)
    table_stats={}
    for m in ranked_methods:
        v=agg({d:e[m] for d,e in complete_means.items()},
              {d:e['single_l2'] for d,e in complete_means.items()})
        v['class_ratio']=stats[m]['class_ratio']
        v['class_datasets']=PDE_COUNT
        v['elo']=elo['ratings'][m]['elo']
        v['rank']=elo['ratings'][m]['rank']
        table_stats[m]=v
    dump(ROOT/'data/table2_ranking.json',dict(passed=True,statistics=table_stats,
         errors=complete_means,elo=elo,reference='single_l2',
         class_scope='21 PDE datasets',dataset_and_elo_scope='22 datasets including ERA5'))
    rows=[]
    for m in elo['ranking']:
        v=table_stats[m]
        if m in MODEL_INFO:
            ident,name,cite,_=MODEL_INFO[m];name=ident+' '+name
        else:
            name,cite=RULES[m]
        rows.append([name,r'\citep{'+cite+'}' if cite else 'Fitting set selection',
                     f"{v['class_ratio']:.3f}",f"{v['dataset_ratio']:.3f}",f"{v['elo']:.0f}"])
    table('paper_baselines_main.tex','llrrr',
          'Model / rule & Component or idea source & PDE class ratio & Dataset ratio & Elo',rows)
    extra=[m for m,v in MODEL_INFO.items() if v[0].startswith('B')]
    # Keep the split exports in the companion package and typeset one combined table.
    for filename,cols in [('paper_baselines_detail_1.tex',extra[:7]),
                          ('paper_baselines_detail_2.tex',extra[7:]),
                          ('paper_baselines_detail.tex',extra)]:
        rows=[]
        for d in order:
            rows.append([tex(NAMES[d])+(r'$^{\dagger}$' if d in EXCLUSIONS else ''),*[fmt(means[d][m]) if m in means[d] else '---' for m in cols],fmt(means[d]['inverse_mse']),fmt(means[d]['full'])])
        rows.append([r'\midrule ERA5',
                     *[fmt(era_baselines[m]) if m in era_baselines else '---' for m in cols],
                     fmt(era['b5__inverse_mse']),fmt(era['b5__full'])])
        table(filename,'l'+'r'*(len(cols)+2),
            'Dataset & '+' & '.join([MODEL_INFO[m][0] for m in cols]+['Inverse','Fitted']),rows)
    # The appendix credits the combined components of each library member.
    # Table 2 keeps compact representative citations and refers to these details.
    library_sources={
        'M1':'li2021fno,perez2018,lyu2023',
        'M2':'li2021fno,perez2018,herde2024',
        'M3':'liu2022convnext,ronneberger2015,perez2018,lyu2023',
        'M4':'yu2026fire,li2021fno',
        'M5':'tripura2022wno,perez2018,lyu2023',
        'M6':'lu2021,lu2022mf',
        'M7':'higdon2008',
        'M8':'wu2024transolver,perez2018',
        'M9':'liu2022convnext,ronneberger2015,perez2018',
    }
    rows=[[v[0],v[1],r'\citep{'+library_sources.get(v[0],v[2])+'}' if v[2] else 'Reference / control',v[3]] for m,v in MODEL_INFO.items() if not v[0].startswith('C')]
    table('baseline_provenance.tex',r'p{.035\linewidth}p{.22\linewidth}p{.25\linewidth}p{.39\linewidth}',
          'ID & Implementation & Source & Adaptation and inclusion rationale',rows)
    for part,chunk in enumerate([rows[:9],rows[9:17],rows[17:]],1):
        table(f'baseline_provenance_{part}.tex',r'p{.045\linewidth}p{.20\linewidth}p{.23\linewidth}p{.39\linewidth}',
              'ID & Implementation & Source & Adaptation and inclusion rationale',chunk)
    result=dict(passed=True,datasets=PDE_COUNT,additional_methods=13,matched_model_dataset_cells=len(proofs),
        partition_means_verified=len(proofs)*3,missing_cells=missing,proof_counts={p:sum(v['proof']==p for v in proofs) for p in sorted({v['proof'] for v in proofs})},
        training_seed=42,evaluation_partitions=[71,172,273],selection_rule=base['rule'],
        scalar_errors_precede_float16_export=True,training_budgets_matched=False,statistics=stats)
    dump(ROOT/'qa/baseline_match_checks.json',result)
    dump(ROOT/'data/matched_baseline_summary.json',dict(means=means,stats=stats,provenance=proofs,missing=missing))
    wb=load_workbook(ROOT/'data/paper_data.xlsx')
    for name,header,rows in [
        ('Matched baselines',['dataset','method','base_seed','relative_l2','percent'],[[d,m,42,e,100*e] for d in order for m,e in means[d].items()]),
        ('Baseline coverage',list(proofs[0]),[list(v.values()) for v in proofs]),
        ('Baseline aggregates',['method','datasets','class_ratio','dataset_ratio','wins','losses'],[[m,v['n'],v['class_ratio'],v['dataset_ratio'],v['wins'],v['losses']] for m,v in stats.items()]),
        ('Headline comparisons',['comparison','datasets','class_ratio','dataset_ratio','wins','losses'],[[m,v['n'],v['class_ratio'],v['dataset_ratio'],v['wins'],v['losses']] for m,v in headlines.items()]),
        ('Table 2 ranking',['method','rank','PDE_class_ratio','dataset_ratio_22','elo','elo_order_std','wins','losses','ties'],
         [[m,table_stats[m]['rank'],table_stats[m]['class_ratio'],table_stats[m]['dataset_ratio'],
           table_stats[m]['elo'],elo['ratings'][m]['order_std'],elo['ratings'][m]['wins'],
           elo['ratings'][m]['losses'],elo['ratings'][m]['ties']] for m in elo['ranking']])]:
        if name in wb:del wb[name]
        ws=wb.create_sheet(name);ws.append(header)
        for row in rows:ws.append(row)
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
    wb.save(ROOT/'data/paper_data.xlsx')
    print('Matched',len(proofs),'additional model/dataset cells; unavailable:',len(missing))
    for m in main_models+['inverse_mse','full']:print(m,stats[m])

if __name__=='__main__':main()
