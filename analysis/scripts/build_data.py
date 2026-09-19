"""Build paper tables and plot data from frozen local experiment snapshots.

Run once with --refresh-snapshots to copy source summaries from the shared workspace.
Subsequent runs use only the bundled data. No model training or test-time selection.
"""
from pathlib import Path
import argparse, csv, hashlib, json, math, re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from paper_scope import retained, filtered_summary, HISTORICAL_COUNT, PDE_COUNT, PAPER_COUNT, EXTENSION_DATASETS

ROOT = Path(__file__).resolve().parents[1]
WORK = ROOT.parent
CLASSES = {
 'elliptic': ['poisson_generated','poisson_generated_v2','poisson_local','darcy_generated','ext__helmholtz_2d','sharp__helmholtz_2d','ext__eikonal_2d','ext__pressure_poisson_poiseuille'],
 'diffusion': ['heat_generated','heat_local'],
 'shocks': ['sharp__burgers_2d','sharp__euler'],
 'convection': ['fluid','lid_driven_cavity_generated','lid_driven_cavity_v2','ext__rayleigh_benard_2d'],
 'porous': ['sharp__porous_medium_2d'],
 'reaction_diffusion': ['ext__cahn_hilliard_2d','sharp__cahn_hilliard','sharp__fisher_kpp_2d','sharp__allen_cahn_2d','sharp__phase_field_crystal_2d'],
 'waves': ['ext__wave_2d','sharp__shallow_water_2d'], 'climate':['era5']}
CLASS_OF = {d:c for c,ds in CLASSES.items() for d in ds}
NAMES = {
 'poisson_generated':'Excluded original Poisson', 'poisson_generated_v2':'Poisson I', 'poisson_local':'Poisson II',
 'darcy_generated':'Darcy', 'ext__helmholtz_2d':'Excluded resonant Helmholtz', 'sharp__helmholtz_2d':'Helmholtz',
 'ext__eikonal_2d':'Eikonal', 'ext__pressure_poisson_poiseuille':'Pressure Poisson',
 'heat_generated':'Heat I', 'heat_local':'Heat II', 'sharp__burgers_2d':'Burgers', 'sharp__euler':'Euler',
 'fluid':'Navier Stokes', 'lid_driven_cavity_generated':'Excluded small cavity', 'lid_driven_cavity_v2':'Cavity',
 'ext__rayleigh_benard_2d':'Rayleigh Bénard', 'sharp__porous_medium_2d':'Porous medium',
 'ext__cahn_hilliard_2d':'Cahn Hilliard I', 'sharp__cahn_hilliard':'Cahn Hilliard II',
 'sharp__fisher_kpp_2d':'Fisher KPP', 'sharp__allen_cahn_2d':'Allen Cahn', 'sharp__phase_field_crystal_2d':'Phase field crystal',
 'ext__wave_2d':'Wave', 'sharp__shallow_water_2d':'Shallow water', 'era5':'ERA5'}
CLASS_ORDER = ('elliptic', 'diffusion', 'convection', 'shocks', 'porous',
               'reaction_diffusion', 'waves', 'climate')
CLASS_LABELS = {'elliptic':'Elliptic', 'diffusion':'Diffusion',
                'convection':'Convection', 'shocks':'Shocks', 'porous':'Porous flow',
                'reaction_diffusion':'Reaction diffusion', 'waves':'Waves',
                'climate':'Climate'}

def ordered_datasets(datasets):
    """One reporting order: PDE class, then display name, with climate last."""
    return sorted(datasets, key=lambda d: (CLASS_ORDER.index(CLASS_OF[d]), NAMES[d]))

DATASET_CITES = {'poisson_local':'niu2024','heat_local':'niu2024','fluid':'niu2024','era5':'hersbach2020,niu2024'}
def cited_name(d):
    return tex(NAMES[d])+(r' \citep{'+DATASET_CITES[d]+'}' if d in DATASET_CITES else '')

EXCLUSIONS = {
 'poisson_local':'inherited numerical defects',
 'ext__cahn_hilliard_2d':'incomplete input conditioning',
 'ext__pressure_poisson_poiseuille':'synthetic low-fidelity fields'}
METHODS = {'single_l2':'Selected model','uniform':'Uniform (all nine)','norm_top3_uniform':'Top three uniform',
 'aligned_norms':'Aligned error control','norm_only_selected':'Norm only selector','best_pair':'Fitted pair',
 'inverse_mse':'Inverse error mixture','full':'Fitted mixture','allpairs':'All pairs FNO'}
from rule_labels import canonical_label, header_label

def digest(p): return hashlib.sha256(p.read_bytes()).hexdigest()
def dump(p,x): p.write_text(json.dumps(x,indent=2,allow_nan=False)+'\n')
def load(p): return json.loads(p.read_text())
def tex(s): return s.replace('&',r'\&').replace('_',r'\_')
def table_label(value):
    value=str(value)
    if value=='---':return r'\textemdash{}'
    value=canonical_label(value)
    value=re.sub(r'(?<=[A-Za-z])--?(?=[A-Za-z])',' ',value)
    value=value.replace(';','.').replace('–',' ')
    return value

def table(path,columns,header,rows):
    header=' & '.join(header_label(table_label(cell)) for cell in header.split(' & '))
    rows=[[table_label(v) for v in row] for row in rows]
    (ROOT/'tables'/path).write_text('\n'.join([r'\begin{tabular}{'+columns+'}',r'\toprule',header+r' \\',r'\midrule']+
      [' & '.join(map(str,r))+r' \\' for r in rows]+[r'\bottomrule',r'\end{tabular}'])+'\n')

def refresh():
    sources={
      'mechanism_summary.json':'mf_field_fieldgate_mechanism_20260913/results/summary.json',
      'fresh_heat_summary.json':'mf_field_fieldgate_fresh_heat_verified_20260913/results/summary.json',
      'feature_findings.md':'mf_field_fieldgate_features_20260913/results/FINDINGS.md',
      'audit_findings.md':'mf_field_ensemble_audit_20260913/FINDINGS.md',
      'dataset_issues.md':'mf_field_session_20260905/datafix/DATASET_ISSUES_2026-09-06.md',
      'repair_protocol.md':'mf_field_automl_repair_20260913/README.md',
      'ensemble_rules.py':'mf_field_automl_repair_20260913/ensemble_rules.py',
      'era5_summary.json':'mf_field_era5_results_for_paper_20260916/nine/results/era5__summary.json',
      'era5_reported.csv':'mf_field_era5_results_for_paper_20260916/nine/results/era5.csv'}
    era_root=WORK/'mf_field_era5_results_for_paper_20260916/nine/results/era5'
    for p in sorted(era_root.glob('*/*')):
        if p.is_file():sources['era5/'+str(p.relative_to(era_root))]=str(p.relative_to(WORK))
    for p in sorted((WORK/'mf_field_fieldgate_mechanism_20260913/runs').glob('*/*')):
        if p.is_file():sources['historical_runs/'+str(p.relative_to(WORK/'mf_field_fieldgate_mechanism_20260913/runs'))]=str(p.relative_to(WORK))
    manifest=[]
    for dest,src in sources.items():
        p=WORK/src;out=ROOT/'data'/dest;out.parent.mkdir(parents=True,exist_ok=True);out.write_bytes(p.read_bytes())
        manifest.append(dict(snapshot=dest,source=src,sha256=digest(p)))
    audit=load(WORK/'mf_field_ensemble_audit_20260913/raw_input_audit.json')
    rows=[]
    for a in audit['rows']:
        counts={int(Path(f['path']).stem.split('l')[-1]):f['shape'][0] for f in a['training_inputs']}
        grids=[]
        for p in sorted((WORK/'mf_field_fieldgate_matched_20260912/export_results').glob('*__'+a['dataset']+'.json')):
            grid=load(p).get('work_grid')
            if grid:grids.append(grid)
        eval_counts=[]
        for p in (WORK/'mf_field_fieldgate_mechanism_20260913/runs').glob(a['dataset']+'__s*/result.json'):
            eval_counts.append(load(p)['budgets'][0]['n_evaluation'])
        rows.append(dict(dataset=a['dataset'],input_dimension=a['input_dimension'],levels=len(counts),
          hf_train_available=counts[a['hf_level']],lf_train_min=min(v for k,v in counts.items() if k!=a['hf_level']),
          lf_train_max=max(v for k,v in counts.items() if k!=a['hf_level']),n_historical_query=a['n_hf_test'],
          lf_overlap=len(a['lf_overlap_rows']),hf_overlap=len(a['hf_overlap_rows']),
          historical_nine_experts=a['in_nine_model_corpus'],work_grid=grids[0] if grids else None,
          evaluation_cases_per_partition=eval_counts[0] if eval_counts else None))
    dump(ROOT/'data'/'dataset_inventory.json',rows)
    manifest.append(dict(snapshot='dataset_inventory.json',source='mf_field_ensemble_audit_20260913/raw_input_audit.json + export metadata + mechanism task counts',
                         source_audit_sha256=digest(WORK/'mf_field_ensemble_audit_20260913/raw_input_audit.json')))
    dump(ROOT/'data'/'snapshot_manifest.json',manifest)

def build_era5():
    """Verify the nine model study and export its complete results."""
    from build_era5_completion import build
    build()
    s=load(ROOT/'data/era5_summary.json');archive=load(ROOT/'data/era5_archive_checks.json')
    assert s['complete'] and s['base_seed']==42 and s['n_base_hf_train']==55
    assert archive['source_manifest_valid'] and archive['summary_sha256']==digest(ROOT/'data/era5_summary.json')
    assert archive['pod_gp_query_rows']==17 and archive['pod_gp_max_query_variation_relative_to_first']==0
    models=s['models'];means=s['means'];by_method={};cases=[];selections=[];weights=[];case_arrays={}
    for seed in [71,172,273]:
        d=ROOT/'data/era5'/f's{seed}';fit=load(d/'fit.json');ev=load(d/'evaluation.json')
        assert digest(d/'locked_weights.npz')==fit['weights_sha256']
        assert digest(d/'fit.json')==ev['fit_sha256']
        assert digest(d/'per_case_errors.npz')==ev['per_case_errors_sha256']
        assert fit['source_manifest_sha256']==archive['source_manifest_sha256']
        assert fit['prediction_sha256']==archive['prediction_hashes']
        assert not fit['final_answers_loaded'] and ev['weights_locked_before_evaluation'] and ev['passed']
        assert ev['n_evaluation']==7 and fit['models']==models
        assert not set(fit['roles']['calibration_pool']) & set(fit['roles']['evaluation'])
        with np.load(d/'per_case_errors.npz') as z:
            assert z['query_rows'].tolist()==list(range(10,17))==fit['roles']['evaluation']
            for method in z.files:
                if method=='query_rows':continue
                values=z[method];assert values.shape==(7,) and np.isfinite(values).all()
                by_method.setdefault(method,[]).append(float(values.mean()))
                case_arrays.setdefault(method,[]).append(values.copy())
                for row,value in zip(z['query_rows'],values):
                    cases.append(dict(seed=seed,query_row=int(row),evaluation_case=int(row)-9,method=method,relative_l2=float(value),error_percent=100*float(value)))
        with np.load(d/'locked_weights.npz') as z:
            for key in z.files:
                w=z[key];assert np.isfinite(w).all() and w.min()>=0 and abs(w.sum()-1)<1e-10
                budget,rule=key.split('__',1)
                for model,value in zip(models,w):weights.append(dict(seed=seed,budget=int(budget[1:]),rule=rule,model=model,weight=float(value)))
            for b in ['5','10']:
                pick=fit['selection'][b];assert pick['chosen']=='full'
                assert np.array_equal(z[f'b{b}__auto'],z[f'b{b}__full'])
                if b=='10':assert fit['roles']['budgets'][b]==list(range(10))
                else:assert len(fit['roles']['budgets'][b])==5
                selections.append(dict(seed=seed,budget=int(b),chosen=pick['chosen'],**{k:100*v for k,v in pick['leave_one_out_scores'].items()}))
    assert set(by_method)==set(means)
    for method,values in by_method.items():assert np.isclose(np.mean(values),means[method],rtol=1e-13,atol=1e-15),method
    with (ROOT/'data/era5_reported.csv').open() as f:
        records=list(csv.DictReader(f))
    assert len(records)==3*len(means)
    for row in records:
        assert int(row['n_evaluation'])==7 and int(row['n_base_hf_train'])==55
        assert np.isclose(float(row['rel_l2']),by_method[row['method']][[71,172,273].index(int(row['seed']))],rtol=1e-13,atol=1e-15)
    def emit_csv(name,rows):
        with (ROOT/'data'/name).open('w',newline='') as f:
            w=csv.DictWriter(f,fieldnames=list(rows[0]));w.writeheader();w.writerows(rows)
    emit_csv('era5_per_case_errors.csv',cases)
    main_methods=[('Inverse error mixture','inverse_mse'),('Fitted mixture','full'),('Selected model (POD GP)','selected_single'),('Training mean field','training_mean_control'),('Uniform mixture','uniform'),('All pairs FNO','allpairs'),('Fine only FNO','hf_only_control')]
    rows=[]
    for label,method in main_methods:
        rows.append([label,*[f"{100*means[method if method.endswith('_control') else 'b'+b+'__'+method]:.4f}" for b in ['5','10']]])
    table('era5_main.tex','lrr',r'Method & $K=5$ & $K=10$',rows)
    names=['FiLM FNO transfer','All pairs FNO','ConvNeXt transfer','Distribution FNO','Wavelet transfer','Retrieved field DeepONet','POD GP','Slice attention corrector','ConvNeXt corrector']
    rows=[[name,f"{100*means['expert:'+model]:.4f}"] for name,model in zip(names,models)]
    rows.extend([[label,f'{100*means[key]:.4f}'] for label,key in [('Fine only FNO control','hf_only_control'),('Training mean field','training_mean_control')]])
    table('era5_experts.tex','lr',r'Model & Error (\%)',rows)
    table('era5_selection.tex','rrlrrr',r'Subset & $K$ & Chosen & Single & Inverse & Full',[[[71,172,273].index(v['seed'])+1,v['budget'],'Fitted' if v['chosen']=='full' else v['chosen'],*[f"{v[k]:.4f}" for k in ['selected_single','inverse_mse','full']]] for v in selections])
    baseline=np.mean(case_arrays['training_mean_control'],axis=0)
    fig,ax=plt.subplots(figsize=(5.8,2.7))
    for method,offset,color,label in [('b5__inverse_mse',-.12,'#E07A32','Inverse error mixture'),('b5__full',.12,'#156082','Fitted mixture')]:
        change=100*(np.mean(case_arrays[method],axis=0)-baseline)
        ax.scatter(change,np.arange(1,8)+offset,s=22,color=color,label=label,zorder=3)
    ax.axvline(0,color='.4',ls='--',lw=.8);ax.set_yticks(range(1,8));ax.invert_yaxis()
    ax.set_ylabel('Evaluation case');ax.set_xlabel('Error change from mean field (percentage points)')
    ax.set_title('ERA5, five fitting examples',loc='left',fontsize=10)
    ax.spines[['top','right']].set_visible(False)
    ax.legend(loc='lower right',fontsize=7,framealpha=1,facecolor='white',edgecolor='none')
    fig.tight_layout();fig.savefig(ROOT/'figures/era5_case_errors.pdf',bbox_inches='tight');fig.savefig(ROOT/'figures/era5_case_errors.png',dpi=180,bbox_inches='tight');plt.close(fig)
    checks=dict(passed=True,means_recomputed=len(means),reported_csv_rows_verified=len(records),partition_hash_checks=3,n_evaluation=7,base_seed=42,n_base_hf=55,models=len(models),auto_selected_full_all_partitions=True,ten_case_budget_reuses_same_pool=True,pod_gp_query_variation=archive['pod_gp_max_query_variation_relative_to_first'],inverse5_gain_vs_mean_percent=100*(1-means['b5__inverse_mse']/means['training_mean_control']),auto5_change_vs_mean_percent=100*(means['b5__auto']/means['training_mean_control']-1))
    dump(ROOT/'qa/era5_checks.json',checks)
    return dict(summary=s,cases=cases,selections=selections,weights=weights,checks=checks)


def build_individual_comparison(summary):
    """Rescore the archived nine experts on the mixture evaluation rows."""
    from collections import Counter
    from matplotlib.colors import TwoSlopeNorm
    datasets=ordered_datasets(summary['5']['mean_errors'])
    from extension_results import partitions as extension_partitions
    means={};flat=[];partitions=[];models=load(ROOT/'data/four_dataset_completion/PLAN.json')['models']
    for ds in datasets:
        per_model=[]
        if ds in EXTENSION_DATASETS:
            added=extension_partitions(ds);partitions.extend(added);per_model=[v['errors'] for v in added]
        for seed in ([] if ds in EXTENSION_DATASETS else [71,172,273]):
            d=ROOT/'data/historical_runs'/f'{ds}__s{seed}'
            fit=load(d/'fit.json');result=load(d/'result.json');split=load(d/'split.json')
            if models is None:models=fit['models']
            assert fit['models']==models and len(models)==9
            assert digest(d/'fit.json')==result['fit_sha256']
            assert digest(d/'locked_weights.npz')==fit['locked_weights_sha256']==result['locked_weights_sha256']
            assert digest(d/'split.json')==fit['split_sha256']
            assert digest(d/'per_case_errors.npz')==result['per_case_errors_sha256']
            assert fit['all_budgets_locked_before_evaluation'] and fit['baseline_replay_passed']
            with np.load(d/'per_case_errors.npz') as z, np.load(d/'locked_weights.npz') as w:
                assert z['row'].tolist()==[v['row'] for v in split['evaluation']]
                assert len(z['row'])==fit['n_evaluation']
                for entries in split['calibration'].values():
                    assert not set(v['group'] for v in split['evaluation']) & set(v['group'] for v in entries)
                errors=np.stack([z['expert:'+m] for m in models],axis=1)
                assert errors.shape==(len(z['row']),9) and np.isfinite(errors).all() and (errors>=0).all()
                assert np.allclose(z['b5__allpairs'],errors[:,models.index('mf_fno_allpairs')],rtol=1e-12,atol=1e-15)
                selected=w['b5__single_l2'];assert np.count_nonzero(selected)==1 and selected.sum()==1
                assert np.allclose(z['b5__single_l2'],errors[:,int(np.argmax(selected))],rtol=1e-12,atol=1e-15)
                for b in result['budgets']:
                    for method,expected in b['errors'].items():
                        assert np.isclose(z[f"b{b['budget']}__{method}"].mean(),expected,rtol=1e-12,atol=1e-15)
                vals={m:float(errors[:,i].mean()) for i,m in enumerate(models)}
                vals.update({m:float(z['b5__'+m].mean()) for m in ['single_l2','inverse_mse','full']})
                per_model.append(vals)
                partitions.append(dict(dataset=ds,seed=seed,n_evaluation=len(z['row']),errors=vals))
        means[ds]={m:float(np.mean([r[m] for r in per_model])) for m in per_model[0]}
        for m in ['single_l2','inverse_mse','full']:
            assert np.isclose(means[ds][m],summary['5']['mean_errors'][ds][m],rtol=1e-12,atol=1e-15)
        for m,e in means[ds].items():
            flat.append(dict(dataset=ds,pde_class=CLASS_OF[ds],method=m,relative_l2=e,error_percent=100*e,ratio_to_calibrated_single=e/means[ds]['single_l2']))
    with (ROOT/'data/individual_model_errors.csv').open('w',newline='') as f:
        out=csv.DictWriter(f,fieldnames=list(flat[0]));out.writeheader();out.writerows(flat)
    labels=['FNO transfer','All pairs FNO','ConvNeXt transfer','Distribution FNO','Wavelet transfer','Retrieved field DeepONet','POD GP','Slice attention corrector','ConvNeXt corrector']
    short=['FiLM FNO transfer','All pairs FNO','ConvNeXt transfer','Distribution FNO','Wavelet transfer','DeepONet','POD GP','Slice attn. corr.','ConvNeXt corr.','Inverse error mixture','Fitted mixture']
    plotted=models+['inverse_mse','full']
    values=np.array([[means[d][m]/means[d]['single_l2'] for m in plotted] for d in datasets])
    fig,ax=plt.subplots(figsize=(6.8,5.0))
    logvalues=np.log2(values)
    im=ax.imshow(np.clip(logvalues,-1,3),aspect='auto',cmap='RdBu_r',norm=TwoSlopeNorm(vmin=-1,vcenter=0,vmax=3))
    ax.set_xticks(range(len(plotted)),short,rotation=50,ha='right',fontsize=7)
    ax.set_yticks(range(len(datasets)),[NAMES[d].replace('--','–')+(' †' if d in EXCLUSIONS else '') for d in datasets],fontsize=7)
    ax.tick_params(length=0);ax.axvline(8.5,color='black',lw=1)
    winners=[int(np.argmin([means[d][m] for m in models])) for d in datasets]
    ax.scatter(winners,range(len(datasets)),s=10,facecolors='none',edgecolors='black',linewidths=.7)
    for i in range(1,len(datasets)):
        if CLASS_OF[datasets[i]]!=CLASS_OF[datasets[i-1]]:ax.axhline(i-.5,color='white',lw=1.2)
    bar=fig.colorbar(im,ax=ax,fraction=.035,pad=.025,extend='both')
    bar.set_ticks([-1,0,1,2,3],labels=['0.5','1','2','4','8']);bar.ax.tick_params(labelsize=7)
    bar.set_label('Error / selected model',fontsize=8)
    fig.tight_layout(pad=.5)
    fig.savefig(ROOT/'figures/individual_models.pdf',bbox_inches='tight')
    fig.savefig(ROOT/'figures/individual_models.png',dpi=200,bbox_inches='tight');plt.close(fig)
    for filename,indices,head in [('individual_direct.tex',range(0,5),['FNO tr.','All pairs','ConvNeXt tr.','Distribution','Wavelet tr.']),('individual_remaining.tex',range(5,9),['DeepONet','POD GP','Slice attn. corr.','ConvNeXt corr.'])]:
        rows=[]
        for d in datasets:
            name=tex(NAMES[d])+(r'$^{\dagger}$' if d in EXCLUSIONS else '')
            rows.append([name,*[f'{100*means[d][models[i]]:.5g}' for i in indices]])
        table(filename,'l'+'r'*len(head),'Dataset & '+' & '.join(head),rows)
    counts=Counter(models[i] for i in winners)
    checks=dict(passed=True,datasets=len(datasets),models=models,partitions_verified=len(partitions),expert_means_recomputed=len(datasets)*len(models),ensemble_means_reproduced=len(datasets)*3,best_expert_counts=dict(counts),distinct_best_experts=len(counts),hindsight_only=True,clipped_heatmap_cells=int(((logvalues<-1)|(logvalues>3)).sum()))
    dump(ROOT/'qa/individual_model_checks.json',checks)
    dump(ROOT/'data/individual_model_summary.json',dict(means=means,partitions=partitions,checks=checks))
    return dict(rows=flat,checks=checks)


def main():
    parser=argparse.ArgumentParser();parser.add_argument('--refresh-snapshots',action='store_true');args=parser.parse_args()
    for d in ['data','tables','figures','qa']:(ROOT/d).mkdir(exist_ok=True)
    if args.refresh_snapshots:refresh()
    from extension_results import extend_primary,update_inventory,CORE_METHODS
    diagnostic=filtered_summary(load(ROOT/'data/mechanism_summary.json'),CLASS_OF)
    dump(ROOT/'data/paper_mechanism_summary.json',diagnostic)
    s=extend_primary(diagnostic,CLASS_OF)
    dump(ROOT/'data/paper_primary_summary.json',s)
    heat=load(ROOT/'data/fresh_heat_summary.json')
    inventory=update_inventory([a for a in load(ROOT/'data/dataset_inventory.json') if retained(a['dataset'])]);primary=s['5'];rows=[];checks=[]
    flat=[]
    for budget,record in s.items():
        for ds in ordered_datasets(record['mean_errors']):
            errors=record['mean_errors'][ds]
            for method,error in errors.items():flat.append(dict(dataset=ds,pde_class=CLASS_OF[ds],calibration_budget=int(budget),method=method,mean_relative_l2=error,error_percent=100*error,ratio_to_selected_single=error/errors['single_l2']))
    with (ROOT/'data/per_dataset_errors.csv').open('w',newline='') as f:
        w=csv.DictWriter(f,fieldnames=list(flat[0]));w.writeheader();w.writerows(flat)
    for m in ['single_l2','allpairs','uniform','inverse_mse','full']:
        c=primary['comparisons']['single_l2'][m]
        rows.append([METHODS[m],f"{c['class_balanced_ratio']:.3f}",f"{c['dataset_balanced_ratio']:.3f}",str(c['improved_over_1pct']),str(c['worsened_over_1pct'])])
    table('main_results.tex','lrrrr','Method & Class balanced & Equal dataset & Wins & Losses',rows)
    rows=[]
    for b in ['5','10','20']:
        comp=s[b]['comparisons']['single_l2']
        rows.append([b,*[f"{comp[m]['class_balanced_ratio']:.4f}" for m in ['inverse_mse','full']],f"{comp['full']['class_balanced_ratio']/comp['inverse_mse']['class_balanced_ratio']:.4f}"])
    table('budgets.tex','rrrr',r'$K$ & Inverse / single & Full / single & Full / inverse',rows)
    rows=[]
    for ds in ordered_datasets(primary['mean_errors']):
        e=primary['mean_errors'][ds];name=tex(NAMES[ds])+(r'$^{\dagger}$' if ds in EXCLUSIONS else '')
        rows.append([name,*[f'{100*e[m]:.6g}' for m in ['single_l2','inverse_mse','full','allpairs']],f"{100*(e['full']/e['single_l2']-1):+.2f}"])
    table('per_dataset.tex','lrrrrr',r'Dataset & Single (\%) & Inverse (\%) & Full (\%) & All pairs (\%) & Change (\%)',rows)
    sens={};sensrows=[]
    for b,record in diagnostic.items():
        means=record['mean_errors'];ds=[d for d in means if d not in EXCLUSIONS];cs=sorted({CLASS_OF[d] for d in ds});v={}
        for m in ['inverse_mse','full','best_pair']:
            lr={d:math.log(means[d][m]/means[d]['single_l2']) for d in ds}
            v[m]={'class_balanced_ratio':math.exp(np.mean([np.mean([lr[d] for d in ds if CLASS_OF[d]==c]) for c in cs])),
                  'dataset_balanced_ratio':math.exp(np.mean(list(lr.values()))),
                  'wins':sum(math.exp(x)<.99 for x in lr.values()),'losses':sum(math.exp(x)>1.01 for x in lr.values())}
        sens[b]=dict(n_datasets=len(ds),n_classes=len(cs),methods=v)
        sensrows.append([b,f"{v['inverse_mse']['class_balanced_ratio']:.4f}",f"{v['full']['class_balanced_ratio']:.4f}",f"{v['full']['dataset_balanced_ratio']:.4f}",v['full']['wins'],v['full']['losses']])
    dump(ROOT/'data/audit_sensitivity.json',dict(exclusions=EXCLUSIONS,results=sens))
    table('audit_sensitivity.tex','rrrrrr',r'$K$ & Inverse (class) & Full (class) & Full (dataset) & Wins & Losses',sensrows)
    rows=[]
    for b in ['5','10','20']:
        e=heat['nominal'][b]['errors']['heat_generated']
        rows.append([b,*[f'{100*e[m]:.6f}' for m in ['single_l2','inverse_mse','full','shrink_loo','allpairs']]])
    table('fresh_heat.tex','rrrrrr',r'$K$ & Single & Inverse & Full & LOO shrinkage & All pairs',rows)
    table('fresh_heat_main.tex','rrrrr',r'$K$ & Selected & Inverse & Fitted & All pairs FNO',
          [[row[i] for i in [0,1,2,3,5]] for row in rows])
    # Roster counts are available training rows, not a claim of identical optimizer exposure.
    rows=[]
    paper_datasets=set(primary['mean_errors'])|{'era5'}
    assert len(paper_datasets)==PAPER_COUNT
    inventory_by_dataset={a['dataset']:dict(a) for a in inventory}
    # Report the actual input-disjoint ERA5 experiment, not the superseded
    # archive's 65 HF / 7 query counts. Preserve original inventory provenance.
    era_plan=load(ROOT/'data/era5_completion/nine/PLAN.json')['datasets']['era5']
    era_lf=[level['kept_train_count'] for fid,level in era_plan['levels'].items()
            if int(fid)!=era_plan['hf_level']]
    inventory_by_dataset['era5'].update(hf_train_available=era_plan['n_base_hf_train'],
                                       n_historical_query=era_plan['n_query'],
                                       lf_train_min=min(era_lf),lf_train_max=max(era_lf))
    dump(ROOT/'qa/roster_training_counts.json',dict(passed=True,
         era5=dict(base_hf=era_plan['n_base_hf_train'],reserved_queries=era_plan['n_query'],
                   minimum_lf=min(era_lf),maximum_lf=max(era_lf)),
         source='data/era5_completion/nine/PLAN.json',original_inventory_preserved=True))
    for dataset in ordered_datasets(paper_datasets):
        a=inventory_by_dataset[dataset]
        name=cited_name(a['dataset']);grid=a['work_grid'];gridtext=r'$'+r'\times'.join(map(str,grid))+'$' if grid else '---'
        state=CLASS_LABELS[CLASS_OF[a['dataset']]]
        rows.append([name,a['input_dimension'],a['levels'],a['lf_train_min'],a['hf_train_available'],a['n_historical_query'],gridtext,state])
    dump(ROOT/'data/paper_roster.json',dict(dataset_count=PAPER_COUNT,historical_aggregate_count=HISTORICAL_COUNT,pde_aggregate_count=PDE_COUNT,
         datasets=ordered_datasets(paper_datasets),
         era5_completed_experts=9,era5_pending_experts=[],
         source='Twenty one PDE datasets plus ERA5. Four completed reaction diffusion datasets extend the retained seventeen. Detailed mechanism and loss controls cover the original seventeen only. Prior exclusions remain unchanged.'))
    table('roster.tex','lrrrrrll',r'Dataset & $d$ & Levels & $N_{L,\min}$ & $N_H$ & Query & Work grid & Class',rows)
    # Independently reproduce every aggregate, to protect against denominator errors.
    for b,record in s.items():
        for m in CORE_METHODS:
            comp=record['comparisons']['single_l2'][m];lr={d:math.log(e[m]/e['single_l2']) for d,e in record['mean_errors'].items()}
            classes=sorted({CLASS_OF[d] for d in lr});recalc=math.exp(np.mean([np.mean([lr[d] for d in lr if CLASS_OF[d]==c]) for c in classes]))
            assert abs(recalc-comp['class_balanced_ratio'])<1e-12,(b,m)
            assert abs(math.exp(np.mean(list(lr.values())))-comp['dataset_balanced_ratio'])<1e-12
            checks.append([b,m,recalc])
    plt.rcParams.update({'font.family':'DejaVu Sans','font.size':8,'axes.spines.top':False,'axes.spines.right':False,'pdf.fonttype':42})
    colors={'full':'#156082','inverse_mse':'#E07A32','best_pair':'#627A44'}
    fig,ax=plt.subplots(1,2,figsize=(6.0,2.7),gridspec_kw={'width_ratios':[1.2,1]})
    cs=[c for c in CLASS_ORDER if c!='climate']
    labels=[CLASS_LABELS[c] for c in cs]
    yy=np.arange(len(cs))
    for m,offset in [('inverse_mse',-.13),('full',.13)]:
        v=[primary['comparisons']['single_l2'][m]['class_ratios'][c] for c in cs]
        ax[0].plot(v,yy+offset,'o',color=colors[m],markersize=4,label=METHODS[m])
    ax[0].axvline(1,color='.45',ls='--',lw=.8);ax[0].set_yticks(yy,labels);ax[0].invert_yaxis();ax[0].set_xlim(.50,1.10)
    ax[0].set_xlabel('Error / selected model');ax[0].set_title('(a) Class level results, K = 5',loc='left',fontsize=9)
    for m in ['inverse_mse','full']:
        vals=[s[b]['comparisons']['single_l2'][m]['class_balanced_ratio'] for b in ['5','10','20']]
        ax[1].plot([5,10,20],vals,'o-',color=colors[m],markersize=4,label=METHODS[m])
    ax[1].axhline(1,color='.45',ls='--',lw=.8);ax[1].set_ylim(.80,1.025);ax[1].set_xticks([5,10,20]);ax[1].set_xlabel('Additional fitting examples')
    ax[1].set_ylabel('Class balanced error ratio');ax[1].set_title('(b) Fitting budget',loc='left',fontsize=9);ax[1].legend(fontsize=6.5,frameon=True,facecolor='white',edgecolor='none',framealpha=1,loc='upper right')
    fig.tight_layout(w_pad=1.5);fig.savefig(ROOT/'figures/results_overview.pdf',bbox_inches='tight');fig.savefig(ROOT/'figures/results_overview.png',dpi=180,bbox_inches='tight');plt.close(fig)
    era=build_era5()
    individual=build_individual_comparison(s)
    # Workbook contains every plotted number and the transformations/source manifest.
    wb=Workbook();ws=wb.active;ws.title='Dataset errors';ws.append(list(flat[0]))
    for r in flat:ws.append(list(r.values()))
    ws=wb.create_sheet('Aggregate ratios');ws.append(['budget','method','class_ratio','dataset_ratio','wins_gt1pct','losses_gt1pct'])
    for b,record in s.items():
        for m in CORE_METHODS:
            c=record['comparisons']['single_l2'][m];ws.append([int(b),m,c['class_balanced_ratio'],c['dataset_balanced_ratio'],c['improved_over_1pct'],c['worsened_over_1pct']])
    ws=wb.create_sheet('Class plot');ws.append(['budget','class','method','ratio_to_single'])
    for m in ['inverse_mse','full']:
        for c in cs:ws.append([5,c,m,primary['comparisons']['single_l2'][m]['class_ratios'][c]])
    ws=wb.create_sheet('Fresh heat');ws.append(['setting','budget','method','relative_l2','percent'])
    for setting,bs in heat.items():
        for b,record in bs.items():
            for m,e in record['errors']['heat_generated'].items():ws.append([setting,int(b),m,e,e*100])
    ws=wb.create_sheet('ERA5 results');ws.append(['method','relative_l2','percent','n_evaluation','n_base_hf'])
    for method,e in era['summary']['means'].items():ws.append([method,e,100*e,7,55])
    for name,rows in [('ERA5 cases',era['cases']),('ERA5 selection',era['selections']),('ERA5 weights',era['weights'])]:
        ws=wb.create_sheet(name);ws.append(list(rows[0]))
        for row in rows:ws.append(list(row.values()))
    for title,file in [('ERA5 baselines','era5_baselines_reported.csv'),('ERA5 baseline cases','era5_baselines_per_case.csv')]:
        ws=wb.create_sheet(title)
        with (ROOT/'data'/file).open() as f:
            for row in csv.reader(f):ws.append(row)
    ws=wb.create_sheet('Individual models');ws.append(list(individual['rows'][0]))
    for row in individual['rows']:ws.append(list(row.values()))
    ws=wb.create_sheet('Sources and definitions');ws.append(['Item','Definition / source'])
    for r in load(ROOT/'data/snapshot_manifest.json'):ws.append([r['snapshot'],r['source']])
    ws.append(['Individual model comparison','Same archived evaluation rows as the 21 PDE mixtures; means across three partitions; circles mark hindsight best individual among nine, never selected using test answers for deployment. Colors show ratios to selected model, clipped outside [0.5,8].'])
    ws.append(['Paper scope','22 datasets: 21 PDE datasets plus ERA5. Reporting scope was narrowed after inspection of historical results. Historical data versions and limitations remain.'])
    ws.append(['ERA5 scope','Nine model experiment; 55 base HF, 5/10 fitting, seven evaluation inputs; not added to the 21 PDE class aggregates. All nine models and their mixtures are complete. All twelve baseline entries represent eleven independent fits, with B10 reusing B9. The separate training mean supplies B13.'])
    ws.append(['ERA5 partitions','K=5 varies fitting subset; K=10 repeats the same ten cases. All share the same seven evaluation cases and base seed42.'])
    ws.append(['ERA5 case plot','Mean over three fitting subsets of per-case relative L2 minus training-mean baseline, times100. Negative percentage-point change improves error. Case labels are indices, not years.'])
    for row in [('Relative L2','Per-case Euclidean residual norm / max(target norm,1e-8); average cases, then partitions.'),('Class ratio','Geometric mean of dataset error ratios within class, then geometric mean over represented classes.'),('Equal dataset ratio','Geometric mean of dataset error ratios; each dataset one vote.'),('Win/loss threshold','ratio <0.99 / >1.01; remaining entries within 1%.'),('Fitting count','Additional to all base-model training labels.'),('Replicates','Three fitting partitions, same base seed42; not independent base-model seeds.'),('Historical status','Development data; results already informed design; no significance claim.'),('Audit sensitivity','Excluded3 by documented data properties; not a certification of remaining generators.')]:ws.append(list(row))
    for ws in wb:
        ws.freeze_panes='A2';ws.auto_filter.ref=ws.dimensions
        for cell in ws[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='156082')
        for col in ws.columns:
            letter=col[0].column_letter;ws.column_dimensions[letter].width=min(70,max(16,max(len(str(c.value or '')) for c in list(col)[:200])+2))
    wb.save(ROOT/'data/paper_data.xlsx')
    assert len(primary['mean_errors'])==PDE_COUNT
    per_partition_total=sum(a['evaluation_cases_per_partition'] or 0 for a in inventory if a['dataset'] in primary['mean_errors'])
    assert per_partition_total==614
    dump(ROOT/'qa/data_checks.json',dict(passed=True,aggregates_recomputed=len(checks),historical_datasets=HISTORICAL_COUNT,pde_datasets=PDE_COUNT,
         audit_subset_datasets=HISTORICAL_COUNT-len(EXCLUSIONS),per_partition_eval_total=per_partition_total,era5_means_recomputed=era['checks']['means_recomputed'],individual_model_means_recomputed=individual['checks']['expert_means_recomputed'],snapshot_hashes={r['snapshot']:digest(ROOT/'data'/r['snapshot']) for r in load(ROOT/'data/snapshot_manifest.json')}))
    print('Built tables, figure, CSV, workbook; verified',len(checks),'aggregate comparisons.')

if __name__=='__main__':main()
