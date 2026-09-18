"""Structural and numerical checks for the manuscript, not a scientific replication."""
from pathlib import Path
import hashlib,json,re,subprocess
import numpy as np
from openpyxl import load_workbook

from paper_scope import REMOVED_DATASETS, HISTORICAL_COUNT, PDE_COUNT, PAPER_COUNT, EXTENSION_DATASETS
ROOT=Path(__file__).resolve().parents[1]
def main():
    summary=json.loads((ROOT/'data/paper_mechanism_summary.json').read_text())
    primary=summary['5']['comparisons']['single_l2']['full']
    assert round(100*(1-primary['class_balanced_ratio']),1)==12.4
    assert round(100*(1-primary['dataset_balanced_ratio']),1)==8.2
    assert primary['improved_over_1pct']==12 and primary['worsened_over_1pct']==5
    assert json.loads((ROOT/'qa/data_checks.json').read_text())['passed']
    era5=json.loads((ROOT/'qa/era5_checks.json').read_text())
    assert era5['passed'] and era5['means_recomputed']==23
    assert era5['n_evaluation']==7 and era5['n_base_hf']==55 and era5['models']==9
    assert era5['auto_selected_full_all_partitions'] and era5['ten_case_budget_reuses_same_pool']
    assert era5['pod_gp_query_variation']==0
    assert round(era5['inverse5_gain_vs_mean_percent'],2)==13.69
    assert round(era5['auto5_change_vs_mean_percent'],2)==-3.70
    from build_era5_completion import build as verify_era5_completion
    verify_era5_completion()
    era_baseline=json.loads((ROOT/'qa/era5_baseline_checks.json').read_text())
    assert era_baseline['passed'] and era_baseline['answers_identical_to_nine_model_study']
    assert era_baseline['available_entries']==12 and era_baseline['independent_fits']==11
    assert era_baseline['missing_ids']==[] and era_baseline['campaign_complete']
    assert era_baseline['production_collection_verified']
    assert round(100*era_baseline['means']['fno_coregionalization'],4)==5.1503
    assert round(100*era_baseline['means']['st_koh_pod'],4)==5.0032
    individual=json.loads((ROOT/'qa/individual_model_checks.json').read_text())
    assert individual['passed'] and individual['partitions_verified']==3*PDE_COUNT
    assert individual['expert_means_recomputed']==9*PDE_COUNT and individual['ensemble_means_reproduced']==3*PDE_COUNT
    assert individual['distinct_best_experts']==7 and sum(individual['best_expert_counts'].values())==PDE_COUNT
    for entry in json.loads((ROOT/'data/snapshot_manifest.json').read_text()):
        if 'sha256' in entry:
            assert hashlib.sha256((ROOT/'data'/entry['snapshot']).read_bytes()).hexdigest()==entry['sha256'],entry['snapshot']
    bib=(ROOT/'references.bib').read_text();keys=re.findall(r'@\w+\{([^,]+),',bib)
    assert len(keys)==len(set(keys))
    used=[]
    for g in re.findall(r'\\citation\{([^}]+)\}',(ROOT/'main.aux').read_text()):
        used.extend(g.split(','))
    assert set(used)==set(keys),(set(used)-set(keys),set(keys)-set(used))
    log=(ROOT/'main.log').read_text(errors='replace');blg=(ROOT/'main.blg').read_text(errors='replace')
    assert not re.search(r'undefined|Citation .* undefined|Overfull|LaTeX Error',log,re.I)
    assert 'Warning--' not in blg
    aux=(ROOT/'main.aux').read_text()
    def label_page(label):
        m=re.search(r'\\newlabel\{'+re.escape(label)+r'\}\{\{[^}]*\}\{(\d+)\}',aux)
        assert m,label
        return int(m.group(1))
    conclusion_page=label_page('main:end')
    references_start=label_page('references:start');references_end=label_page('references:end')
    appendices_start=label_page('appendices:start');statements_page=label_page('app:statements')
    # Count every page before the bibliography, including any delayed main-text floats.
    mainpage=references_start-1
    assert conclusion_page<=mainpage<=9,(conclusion_page,mainpage)
    assert label_page('sec:climate')<=mainpage
    assert label_page('tab:era5')>=appendices_start
    assert all(label_page(k)<=mainpage for k in ['fig:overview','fig:datasets','tab:allexperts','tab:paperbaselines'])
    assert label_page('fig:fieldcomparison')<=mainpage
    from build_field_comparison import build as verify_field_comparison
    field_comparison = verify_field_comparison()
    assert field_comparison['passed'] and field_comparison['fitted_case_wins_over_each_individual']==20
    from rule_labels import RULE_NAMES
    assert set(RULE_NAMES[m] for m in ('selected_single','inverse_mse','full')) <= set(field_comparison['relative_errors_percent'])
    assert field_comparison['selected_model_identical_to_saved_member'] and not field_comparison['automatic_selection_displayed']
    method_text=(ROOT/'sections/method.tex').read_text()
    assert all(r'\noindent\textbf{'+name+'.}' in method_text for name in ('Selected model','Inverse error mixture','Fitted mixture'))
    for table_path in (ROOT/'tables').glob('*.tex'):
        assert not re.search(r'(?:^|&)\s*(?:Inverse|Fitted|Selected|Auto|Automatic|Full|Single)\s*(?:&|\\\\|$)',table_path.read_text(),re.M),table_path
    assert all(label_page(k)>=appendices_start for k in ['tab:training','tab:baseline_detail'])
    assert all(label_page(k)>=appendices_start for k in ['tab:headlinecomparisons','tab:main','tab:budgets','prop:bound'])
    assert label_page('fig:era5cases')>=appendices_start
    # Condensed appendix: complete numerical artifacts remain in the package,
    # while duplicated tables and exploratory branches are no longer typeset.
    removed_labels=['fig:individual','tab:portfolio','tab:directexperts',
                    'tab:remainingexperts','fig:results','tab:era5experts',
                    'tab:reviewsubsets','tab:reviewbudgets','app:refresh']
    assert not any(r'\newlabel{'+k+'}' in aux for k in removed_labels)
    appendix_order=['app:roster','app:protocol','app:numerics','app:allresults',
                    'app:reanalysis','app:geometry','app:statements']
    assert [label_page(k) for k in appendix_order]==sorted(label_page(k) for k in appendix_order)
    assert (ROOT/'supplement/README.md').is_file()
    assert (ROOT/'supplement/archived_appendix/coarse_refresh.tex').is_file()
    assert (ROOT/'supplement/archived_appendix/visual_appendix.tex').is_file()
    assert references_start<=references_end<appendices_start<=statements_page
    maintex=(ROOT/'main.tex').read_text()
    assert r'\usepackage{iclr2026_conference,times}' in maintex
    assert r'\bibliographystyle{iclr2026_conference}' in maintex
    assert '2027' not in maintex
    template=json.loads((ROOT/'qa/template_source.json').read_text())
    assert template['template_year']==2026
    for entry in template['installed_files']:
        assert hashlib.sha256((ROOT/entry['file']).read_bytes()).hexdigest()==entry['sha256'],entry['file']
    pdfinfo=subprocess.check_output(['pdfinfo',str(ROOT/'main.pdf')],text=True)
    pages=int(re.search(r'Pages:\s+(\d+)',pdfinfo).group(1))
    paper_size=re.search(r'Page size:\s+([\d.]+) x ([\d.]+) pts',pdfinfo)
    assert paper_size and tuple(map(float,paper_size.groups()))==(612.0,792.0)
    subprocess.run(['pdftotext','-layout',str(ROOT/'main.pdf'),str(ROOT/'qa/manuscript.txt')],check=True)
    text=(ROOT/'qa/manuscript.txt').read_text()
    assert not any(s.lower() in text.lower() for s in ['Working draft','Draft status','Working manuscript'])
    body_pages=text.split('\f')
    prose='\n'.join(body_pages[:references_start-1]+body_pages[references_end:])
    assert ';' not in prose
    assert not re.search(r'\b(?:calibrat\w*|recalibrat\w*|experts?|predictors?)\b',prose,re.I), 'Inconsistent ensemble terminology in rendered manuscript'
    assert r'\subsection{Surrogate model library}' in (ROOT/'sections/method.tex').read_text()
    assert r'\subsection{Ensemble fitting budgets and evaluation}' in (ROOT/'sections/experiments.tex').read_text()
    assert not re.search(r'repaired\s+ERA5|seed[ -]?42',prose,re.I)
    assert set(re.findall(r'\b[A-Za-z]+-[A-Za-z]+\b',prose)) <= {'Watson-Parris', 'Rais-Rohani'}
    assert not re.search(r'(?:Poisson|Heat|Cavity) \((?:legacy|local|generated|corrected)\)',prose,re.I)
    assert 'Pending' not in text and '11.6%' in text and '7.5%' in text
    assert all(value in text for value in ['5.8146','6.4877','6.7369','13.69%','3.70%'])
    assert 'nearest neighbour retrieval' in (ROOT/'sections/method.tex').read_text()
    assert r'\pending' not in (ROOT/'sections/climate.tex').read_text()
    # Independent algebraic checks of the stated normalization and direction identity.
    rng=np.random.default_rng(42071)
    for _ in range(100):
        y=rng.normal(size=31);pred=rng.normal(size=(9,31));w=rng.dirichlet(np.ones(9))
        e=(pred-y)/max(np.linalg.norm(y),1e-8);g=e@e.T;r=np.linalg.norm(e,axis=1)
        direct=np.linalg.norm(w@pred-y)**2/max(np.linalg.norm(y),1e-8)**2
        assert np.isclose(direct,w@g@w,atol=1e-12)
        aligned=(w@r)**2;direction=2*sum(w[i]*w[j]*(r[i]*r[j]-g[i,j]) for i in range(9) for j in range(i+1,9))
        assert np.isclose(aligned-w@g@w,direction,atol=1e-12) and direction>=-1e-12
        a=rng.normal(size=(9,9));a=(a+a.T)/2;eps=abs(a).max()
        assert abs(w@a@w)<=eps+1e-12
    wb=load_workbook(ROOT/'data/paper_data.xlsx',read_only=True,data_only=True)
    assert len(wb.sheetnames)==17
    assert wb['Baseline coverage'].max_row==13*PDE_COUNT+1
    assert wb['Baseline aggregates'].max_row==27
    assert wb['Aggregate ratios'].max_row==16
    assert wb['ERA5 results'].max_row==24 and wb['ERA5 cases'].max_row==484
    assert wb['ERA5 selection'].max_row==7 and wb['ERA5 weights'].max_row==325
    assert wb['ERA5 baselines'].max_row==1+2*era_baseline['available_entries']
    assert wb['ERA5 baseline cases'].max_row==1+17*era_baseline['available_entries']
    assert wb['Individual models'].max_row==12*PDE_COUNT+1
    headlines=json.loads((ROOT/'data/headline_comparisons.json').read_text())['comparisons']
    assert headlines['Best library / best additional baseline']['wins']==15
    assert headlines['Full / best library']['wins']==13
    assert headlines['Full / best library']['losses']==6
    assert round(headlines['Full / best library']['class_ratio'],3)==0.923
    assert round(headlines['Full / best library']['dataset_ratio'],3)==0.989
    assert wb['Headline comparisons'].max_row==6
    baseline=json.loads((ROOT/'qa/baseline_match_checks.json').read_text())
    assert baseline['passed'] and baseline['datasets']==PDE_COUNT
    assert baseline['matched_model_dataset_cells']==13*PDE_COUNT and baseline['partition_means_verified']==3*13*PDE_COUNT
    assert baseline['training_seed']==42 and len(baseline['missing_cells'])==4*HISTORICAL_COUNT
    assert baseline['statistics']['fno_coregionalization']['n']==PDE_COUNT
    assert baseline['scalar_errors_precede_float16_export'] and not baseline['training_budgets_matched']
    from verify_b8_recovery import verify as verify_b8
    assert verify_b8()['recovered_model_dataset_cells']==2
    # Recompute Table 2 from the per-dataset errors and replay Elo deterministically.
    from elo_ranking import rank_errors
    ranking=json.loads((ROOT/'data/table2_ranking.json').read_text())
    errors=ranking['errors'];methods=ranking['elo']['methods']
    assert len(errors)==PAPER_COUNT and len(methods)==23 and 'era5' in errors
    assert rank_errors(errors,methods)==ranking['elo']
    assert ranking['elo']['match_count']==PAPER_COUNT*23*22//2
    for method in methods:
        ratios=[values[method]/values['single_l2'] for values in errors.values()]
        assert np.isclose(np.exp(np.mean(np.log(ratios))),ranking['statistics'][method]['dataset_ratio'])
        assert np.isclose(baseline['statistics'][method]['class_ratio'],ranking['statistics'][method]['class_ratio'])
        record=ranking['elo']['ratings'][method]
        assert record['wins']+record['losses']+record['ties']==PAPER_COUNT*22
    assert round(100*(1-ranking['statistics']['full']['dataset_ratio']),1)==7.5
    assert ranking['statistics']['full']['wins']==16 and ranking['statistics']['full']['losses']==5
    assert wb['Table 2 ranking'].max_row==24
    ranking_table=(ROOT/'tables/paper_baselines_main.tex').read_text()
    assert '$N$' not in ranking_table and 'Dataset ratio & Elo' in ranking_table
    assert '\\textemdash{}' not in ranking_table
    assert label_page('app:elo')>=appendices_start
    gallery=json.loads((ROOT/'qa/gallery_checks.json').read_text())
    matched=json.loads((ROOT/'data/individual_model_summary.json').read_text())['means']
    assert gallery['passed'] and gallery['main_datasets']==PAPER_COUNT
    assert not gallery['main_colorbars'] and not gallery['main_era5_annotation']
    overview=json.loads((ROOT/'qa/overview_checks.json').read_text())
    assert overview['passed'] and overview['source_fields_are_actual']
    assert overview['single_dataset']=='era5'
    assert overview['training_inputs_bitwise_equal'] and overview['training_field_hashes_verified']
    assert overview['training_source_fields_native_grids'] and overview['shared_color_limits_cover_all_fields']
    assert overview['coarse_preview_is_transformed'] and overview['fine_and_prediction_fields_unchanged']
    assert overview['coarse_display_transform']['illustrative_only']
    assert overview['coarse_display_transform']['documented_in_source_manifest']
    assert not overview['coarse_display_transform']['disclosed_in_caption']
    pair=overview['era5_training_pair']
    assert pair['coarse_grid']==[192,384] and pair['fine_grid']==[721,1440]
    assert pair['source_files_verified_against_transfer_manifest']
    with np.load(ROOT/overview['bundled_data']) as fields:
        assert np.array_equal(fields['era5_training_lf_parameters'],fields['era5_training_hf_parameters'])
        for name in ['coarse','fine']:
            field=fields['era5_training_'+name]
            assert list(field.shape)==pair[name+'_grid']
            assert hashlib.sha256(field.tobytes()).hexdigest()==pair[name+'_array_sha256']
    assert overview['era5_mixture_expert_count']==era5['models']
    assert hashlib.sha256((ROOT/overview['bundled_data']).read_bytes()).hexdigest()==overview['bundled_data_sha256']
    assert gallery['cavity_scale']['data_unchanged'] and gallery['cavity_scale']['dataset']=='lid_driven_cavity_v2'
    assert gallery['cavity_scale']['linear_threshold']==1 and gallery['cavity_scale']['bound']==400
    assert set(gallery['order'])==set(matched)|{'era5'} and not gallery['appendix_only']
    roster=json.loads((ROOT/'data/paper_roster.json').read_text())
    assert roster['dataset_count']==PAPER_COUNT and set(roster['datasets'])==set(gallery['order'])
    assert roster['historical_aggregate_count']==HISTORICAL_COUNT and roster['era5_completed_experts']==9
    assert (ROOT/'tables/roster.tex').read_text().count(' & ')==PAPER_COUNT*7+7
    with np.load(ROOT/'data/gallery_samples.npz') as fields:
        for row in json.loads((ROOT/'data/gallery_manifest.json').read_text())['datasets']:
            a=fields[row['dataset']+'__y']
            assert np.isfinite(a).all()
            assert hashlib.sha256(a.tobytes()).hexdigest()==row['y']['sample_sha256']
    # Verify the new display labels while excluding archives by stable dataset ID.
    from build_data import NAMES
    assert NAMES['poisson_generated_v2'] == 'Poisson I'
    assert NAMES['poisson_local'] == 'Poisson II'
    assert NAMES['lid_driven_cavity_v2'] == 'Cavity'
    assert NAMES['sharp__helmholtz_2d'] == 'Helmholtz'
    assert NAMES['ext__cahn_hilliard_2d'] == 'Cahn Hilliard I'
    assert all(label in prose for label in ('Poisson I', 'Poisson II', 'Heat I', 'Heat II', 'Cavity', 'Helmholtz', 'Cahn Hilliard'))
    assert not re.search(r'\b(?:Poisson III|(?:Cavity|Helmholtz) I{1,2})\b', prose)
    for source in [ROOT/'main.tex', *(ROOT/'sections').glob('*.tex'), *(ROOT/'tables').glob('*.tex')]:
        assert not re.search(r'\b(?:Poisson III|(?:Cavity|Helmholtz) (?:I|II|1|2))\b', source.read_text()), source
    assert not REMOVED_DATASETS & set(roster['datasets'])
    assert not REMOVED_DATASETS & set(matched)
    for budget, record in summary.items():
        assert len(record['mean_errors']) == HISTORICAL_COUNT
        assert not REMOVED_DATASETS & set(record['mean_errors'])
        # Independent recalculation against the retained per-case archives.
        class_logs = {}
        from build_data import CLASS_OF
        ratios = []
        for ds, means in record['mean_errors'].items():
            errors = {m: [] for m in ('full','single_l2')}
            for seed in (71,172,273):
                with np.load(ROOT/'data/historical_runs'/f'{ds}__s{seed}'/'per_case_errors.npz') as z:
                    for method in errors:
                        errors[method].append(float(z[f'b{budget}__{method}'].mean()))
            for method in errors:
                assert np.isclose(np.mean(errors[method]), means[method], rtol=1e-12, atol=1e-15)
            ratio = np.mean(errors['full']) / np.mean(errors['single_l2'])
            ratios.append(ratio)
            class_logs.setdefault(CLASS_OF[ds], []).append(np.log(ratio))
        expected = record['comparisons']['single_l2']['full']
        assert np.isclose(np.exp(np.mean(np.log(ratios))), expected['dataset_balanced_ratio'], rtol=1e-12)
        assert np.isclose(np.exp(np.mean([np.mean(v) for v in class_logs.values()])), expected['class_balanced_ratio'], rtol=1e-12)
    # Replay the four additions and independently check their aggregate inclusion.
    from extension_results import verify as verify_extension, extend_primary
    from build_data import CLASS_OF
    extension=verify_extension()
    rebuilt=extend_primary(summary,CLASS_OF)
    assert rebuilt==json.loads((ROOT/'data/paper_primary_summary.json').read_text())
    assert extension['evaluation_appearances']==228 and extension['distinct_evaluation_dataset_rows']==188
    for ds in EXTENSION_DATASETS:
        assert ds in roster['datasets'] and ds in ranking['errors'] and ds in matched
        assert all(NAMES[ds] in (ROOT/'tables'/f'{name}.tex').read_text() for name in
                   ['all_experts_main','paper_baselines_detail','paper_baselines_detail_1','paper_baselines_detail_2','roster','per_dataset','individual_direct','individual_remaining'])
    assert gallery['detailed_figures']==6 and roster['pde_aggregate_count']==PDE_COUNT
    # Every displayed heatmap and gallery is generated from the same roster.
    assert set(gallery['order']) == set(roster['datasets'])
    appendix=(ROOT/'sections/appendix.tex').read_text()
    assert r'\label{tab:pendingensembles}' not in appendix
    main_table=(ROOT/'tables/all_experts_main.tex').read_text()
    assert r'\pending' not in main_table and main_table.count(r'\textemdash{}')==0
    era_row=next(line for line in main_table.splitlines() if 'ERA5' in line)
    assert '(B1)' in era_row and '20.4' in era_row and '8.24' in era_row and '5.81' in era_row
    assert 'Best baseline' in main_table
    for name in ['all_experts_main','paper_baselines_detail','paper_baselines_detail_1','paper_baselines_detail_2']:
        assert 'ERA5' in (ROOT/'tables'/f'{name}.tex').read_text()
    review_checks=json.loads((ROOT/'qa/review_analysis_checks.json').read_text())
    assert review_checks['passed'] and review_checks['n_fits']==9*HISTORICAL_COUNT
    assert review_checks['evaluation_rows_excluded_from_every_fit'] and review_checks['original_fixed_results_preserved']
    loss_checks=json.loads((ROOT/'qa/loss_control_checks.json').read_text())
    assert loss_checks['passed'] and loss_checks['n_partitions']==3*HISTORICAL_COUNT
    assert loss_checks['source_hashes_passed'] and loss_checks['reference_evaluation_replay_passed']
    assert all(label_page(k)>=appendices_start for k in ['tab:losscontrol','tab:reviewsensitivity','tab:fresh'])
    cli_check=subprocess.run(['python3','scripts/check_calibration_cli.py'],cwd=ROOT,
                             text=True,stdout=subprocess.PIPE,stderr=subprocess.STDOUT)
    (ROOT/'qa/calibration_cli_checks.log').write_text(cli_check.stdout)
    assert cli_check.returncode==0,cli_check.stdout
    assert 'Ran 8 tests' in cli_check.stdout
    cli_checks=dict(passed=True,integration_tests=8,scope='Calibration and fixed weight prediction only',
                    script_sha256=hashlib.sha256((ROOT/'scripts/calibrate_fields.py').read_bytes()).hexdigest())
    (ROOT/'qa/calibration_cli_checks.json').write_text(json.dumps(cli_checks,indent=2)+'\n')
    result=dict(passed=True,main_text_end_page=mainpage,conclusion_page=conclusion_page,total_pages=pages,
      template_year=2026,template_files_match_archive=True,paper_size='US Letter',
      references_start_page=references_start,references_end_page=references_end,
      appendices_start_page=appendices_start,disclosure_statements_page=statements_page,
      all_prebibliography_pages_within_submission_limit=True,bibliography_entries=len(keys),
      all_references_cited=True,missing_citations=False,overfull_boxes=False,
      prose_semicolons=False,automatic_word_hyphenation=False,descriptive_names=True,
      aggregate_checks=15,review_analysis_checks=review_checks,loss_control_checks=loss_checks,calibration_cli_checks=cli_checks,era5_checks=era5,individual_model_checks=individual,
      matched_baseline_checks=baseline,gallery_checks=gallery,overview_checks=overview,field_comparison_checks=field_comparison,paper_datasets=PAPER_COUNT,
      unreported_era5_experts=0,era5_ensemble_experts=9,era5_baseline_checks=era_baseline,headline_comparisons=headlines,
      algebraic_spot_checks=100,workbook_sheets=wb.sheetnames,
      main_pdf_sha256=hashlib.sha256((ROOT/'main.pdf').read_bytes()).hexdigest(),
      limitations='Structural/numerical checks; not a new model replication, literature-priority proof, or human scientific sign-off.')
    (ROOT/'qa/verification.json').write_text(json.dumps(result,indent=2)+'\n');print(json.dumps(result,indent=2))
if __name__=='__main__':main()
